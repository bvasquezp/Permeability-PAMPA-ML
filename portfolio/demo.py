"""Self-contained synthetic portfolio demo. Python 3.10+; openpyxl optional."""
from __future__ import annotations
import argparse
import csv
import json
from collections import Counter, defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

RAW = [
    ("R001", "2026-09-01", "125.00", "Lab"),
    ("R002", "2026-09-02", "76.50", "Ops"),
    ("R002", "2026-09-02", "76.50", "Ops"),  # duplicate
    ("R003", "invalid", "22.00", "Lab"),      # bad date
    ("R004", "2026-09-05", "oops", "Ops"),     # bad amount
    ("R005", "2026-09-07", "90.20", "Lab"),
]
INVOICES = [("INV01", "120.00"), ("INV02", "80.00"), ("INV03", "60.00")]
PAYMENTS = [("INV01", "P01", "120.00"), ("INV02", "P02", "30.00"),
            ("INV02", "P03", "20.00"), ("OTHER", "P04", "9.00")]


def parse_amount(raw: str) -> Decimal:
    try:
        value = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid amount: {raw!r}") from exc
    if not value.is_finite() or value < 0 or value.as_tuple().exponent < -2:
        raise ValueError(f"Invalid nonnegative currency amount: {raw!r}")
    return value


def write_csv(path: Path, columns: list[str], records: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(records)


def generate(output: Path) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    seen, good, bad = set(), [], []
    for line, (identifier, day, amount, category) in enumerate(RAW, 1):
        try:
            if identifier in seen:
                raise ValueError("duplicate identifier")
            checked_day = date.fromisoformat(day).isoformat()
            checked_amount = parse_amount(amount)
            seen.add(identifier)
            good.append(dict(id=identifier, date=checked_day, amount=f"{checked_amount:.2f}", category=category))
        except ValueError as exc:
            bad.append(dict(source_line=line, id=identifier, reason=str(exc)))
    write_csv(output / "clean_records.csv", ["id", "date", "amount", "category"], good)
    write_csv(output / "rejected_records.csv", ["source_line", "id", "reason"], bad)

    grouped = defaultdict(list)
    for invoice_id, payment_id, amount in PAYMENTS:
        grouped[invoice_id].append((payment_id, parse_amount(amount)))
    rows = []
    for invoice_id, original_due in INVOICES:
        due = parse_amount(original_due)
        paid = sum((v for _, v in grouped.pop(invoice_id, [])), Decimal("0"))
        state = ("matched" if paid == due else "unpaid" if paid == 0 else
                 "partial" if paid < due else "overpaid")
        rows.append(dict(invoice_id=invoice_id, due=f"{due:.2f}",
                         paid=f"{paid:.2f}", balance=f"{due-paid:.2f}", status=state))
    for unknown_id in sorted(grouped):
        rows.append(dict(invoice_id=unknown_id, due="", paid="", balance="", status="unmatched_payment"))
    write_csv(output / "reconciliation.csv", ["invoice_id", "due", "paid", "balance", "status"], rows)

    totals = defaultdict(Decimal)
    for item in good:
        totals[item["category"]] += Decimal(item["amount"])
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        excel = False
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.append(["Category", "Count", "Total"])
        for category in sorted(totals):
            ws.append([category, sum(r["category"] == category for r in good), float(totals[category])])
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 20
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="245646")
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
            row[0].number_format = "#,##0.00"
        wb.save(output / "summary.xlsx")
        excel = True
    report = dict(input_rows=len(RAW), valid_rows=len(good),
                  rejected_rows=len(bad), reconciliation_statuses=dict(Counter(r["status"] for r in rows)),
                  excel_created=excel, note="Synthetic demo only; unknown invoice IDs require manual review.")
    (output / "quality_summary.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("portfolio/output"))
    args = parser.parse_args()
    print(json.dumps(generate(args.output), indent=2))


if __name__ == "__main__":
    main()
